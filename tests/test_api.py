"""Uçtan uca API testi (arayüzün kullandığı uç noktalar, sahte Maps sunucusu ile)."""
import openpyxl
from fastapi.testclient import TestClient

import config
from conftest import write_xlsx
from mock_maps_server import MockMapsServer


def test_label_and_edit_flows_end_to_end(input_dir):
    import main

    with MockMapsServer(latency=(0.05, 0.15)) as srv:
        label_rows = [("Name", "Gmaps")] + [(f"L{i}", srv.url(f"L{i}")) for i in range(6)]
        label_rows.append(("Bad", "https://example.com/nope"))
        lp = write_xlsx(input_dir / "aaa_label.xlsx", {"S": label_rows})
        edit_rows = [("durum", "ad", "link")] + [("X" if i % 2 else None, f"E{i}", srv.url(f"E{i}")) for i in range(10)]
        write_xlsx(input_dir / "cleanpoi_cevre_api.xlsx", {"Mekanlar": edit_rows})

        with TestClient(main.app) as client:
            files = client.get("/api/input/files").json()
            by_name = {f["file"]: f for f in files}
            assert by_name["aaa_label.xlsx"]["schema"]["matches"]["maps_url"]["column"] == "Gmaps"

            # ── Etiketleme akışı
            r = client.post("/api/session/start", json={"start_index": 0})
            assert r.status_code == 200, r.text
            body = r.json()
            assert body["pending"] == 7 and body["columns"]["maps_url"] == "Gmaps"
            served = []
            while True:
                item = client.get("/api/screenshot/next").json()
                if item["status"] == "completed":
                    break
                served.append(item)
                if item["file_url"]:
                    img = client.get(item["file_url"])
                    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"
                    assert client.post("/api/label", json={"row_index": item["row_index"], "category": "Kafe"}).status_code == 200
            assert [s["position"] for s in served] == list(range(7))
            assert served[-1]["error"] and "Geçersiz" in served[-1]["error"]     # hatalı URL raporlandı
            stats = client.get("/api/capture/stats").json()
            assert stats["status"] == "running" and stats["store"]["ok"] >= 6
            assert client.post("/api/session/stop").json()["status"] == "stopped"

            ws = openpyxl.load_workbook(lp).active
            header = [c.value for c in ws[1]]
            col = header.index(config.CATEGORY_COLUMN)
            assert [row[col] for row in ws.iter_rows(min_row=2, max_row=7, values_only=True)] == ["Kafe"] * 6

            # ── Düzenle akışı
            r = client.post("/api/edit/session/start", json={"start_index": 0, "source_file": "cleanpoi_cevre_api.xlsx"})
            assert r.status_code == 200, r.text
            assert r.json()["pending"] == 5 and r.json()["columns"]["status"] == "durum"
            first = client.get("/api/edit/screenshot/next").json()
            assert first["kontrol_value"] == "X" and first["next_name"] == "E3" and first["file_url"]
            assert client.post("/api/edit/action", json={"row_index": first["row_index"], "action": "delete"}).json()["status"] == "deleted"
            assert client.post("/api/edit/action/undo", json={"row_index": first["row_index"]}).json()["undone_action"] == "delete"
            assert client.post("/api/edit/action", json={"row_index": first["row_index"], "action": "delete"}).status_code == 200
            client.post("/api/edit/session/stop")
            exp = client.get("/api/edit/export")
            assert exp.status_code == 200 and "cleanpoi_cevre_api_guncel_" in exp.headers["content-disposition"]
