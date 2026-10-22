"""Etiketleme (DataManager) ve düzenle (EditDataManager) akışı veri testleri."""
import csv
import json
import sqlite3
from pathlib import Path

import openpyxl
import pytest

import config
from conftest import maps_url, write_xlsx
from data_model import ColumnRole


# ── Etiketleme akışı ─────────────────────────────────────────────────────────
def _label_file(input_dir: Path, name="etiket.xlsx", n=30) -> Path:
    rows = [("Rapor başlığı",), ("Mekan Adı", "Harita Bağlantısı", "Not")]
    for i in range(n):
        rows.append((f"Yer {i}", maps_url(i) if i != 5 else None, None))
    return write_xlsx(input_dir / name, {"Liste": rows})


def test_label_flow_detects_columns_and_writes_to_correct_cells(input_dir):
    from data_manager import DataManager
    p = _label_file(input_dir)
    dm = DataManager(p.name)
    assert dm.schema.header_row == 2
    assert dm.schema.column(ColumnRole.MAPS_URL) == "Harita Bağlantısı"
    recs = dm.get_place_records(0)
    assert len(recs) == 29 and 5 not in [r.row_index for r in recs]   # URL'siz satır atlandı
    dm.save_category(3, "Kafe", {"Pasif": "1"})
    dm.save_category(7, "", {"Çıkarılacak Noktalar": "x"})
    dm.close()

    ws = openpyxl.load_workbook(p)["Liste"]
    header = [c.value for c in ws[2]]
    cat, pasif, cik = (header.index(config.CATEGORY_COLUMN) + 1, header.index("Pasif") + 1,
                       header.index("Çıkarılacak Noktalar") + 1)
    assert ws.cell(row=2 + 1 + 3, column=cat).value == "Kafe"
    assert ws.cell(row=2 + 1 + 3, column=pasif).value == "1"
    assert ws.cell(row=2 + 1 + 3, column=1).value == "Yer 3"            # doğru satır
    assert ws.cell(row=2 + 1 + 7, column=cik).value == "x"
    assert ws.cell(row=1, column=1).value == "Rapor başlığı"            # orijinal veri korundu
    assert not list(input_dir.glob("~*"))                              # geçici dosya kalmadı

    dm2 = DataManager(p.name, start_flusher=False)
    assert {3, 7}.isdisjoint(r.row_index for r in dm2.get_place_records(0))
    dm2.close()


def test_label_saves_survive_crash_via_journal(input_dir):
    from data_manager import DataManager
    p = _label_file(input_dir, "crash.xlsx")
    dm = DataManager(p.name, start_flusher=False)     # flush hiç çalışmayacak
    for i in range(10):
        dm.save_category(i, f"K{i}", {})
    # "Çökme": close/flush çağrılmadan yeni örnek açılır
    dm2 = DataManager(p.name, start_flusher=False)
    assert dm2.get_pending_count() == 29 - 9          # 5. satırın URL'si yok, zaten bekleyen değildi
    assert dm2.flush()
    dm2.close()
    ws = openpyxl.load_workbook(p)["Liste"]
    header = [c.value for c in ws[2]]
    col = header.index(config.CATEGORY_COLUMN) + 1
    assert [ws.cell(row=3 + i, column=col).value for i in range(10)] == [f"K{i}" for i in range(10)]


def test_label_locked_file_keeps_entries_in_journal(input_dir, monkeypatch):
    import data_manager as dmod
    p = _label_file(input_dir, "locked.xlsx")
    dm = dmod.DataManager(p.name, start_flusher=False)
    dm.save_category(1, "A", {})

    def locked(*a, **k):
        raise PermissionError("dosya Excel'de açık")
    monkeypatch.setattr(dmod.os, "replace", locked)
    assert dm.flush() is False
    assert len(dm._journal.pending(dm._journal_key)) == 1               # kayıp yok
    monkeypatch.undo()
    assert dm.flush() is True
    assert dm._journal.pending(dm._journal_key) == []
    dm.close()


def test_label_flow_csv(input_dir):
    from data_manager import DataManager
    p = input_dir / "liste.csv"
    with open(p, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["name", "maps"])
        for i in range(5):
            w.writerow([f"Y{i}", maps_url(i)])
    dm = DataManager(p.name)
    dm.save_category(2, "Market", {})
    dm.close()
    rows = list(csv.reader(open(p, encoding="utf-8-sig"), delimiter=";"))
    assert rows[0][:2] == ["name", "maps"] and config.CATEGORY_COLUMN in rows[0]
    assert rows[3][rows[0].index(config.CATEGORY_COLUMN)] == "Market"


# ── Düzenle akışı ────────────────────────────────────────────────────────────
def _edit_file(input_dir: Path, name="cleanpoi_cevre_test.xlsx", n=200) -> Path:
    rows = [("Status", "Business", "Rating", "Map Link", "Seen")]
    import datetime
    for i in range(n):
        rows.append(("REVIEW" if i % 4 == 0 else None, f"Biz {i}", 4.5 if i % 2 else 4,
                     maps_url(i), datetime.datetime(2026, 10, 1, 12, i % 60)))
    rows.append((None, None, None, None, None))
    return write_xlsx(input_dir / name, {"Summary": [("k", "v")], "Rows": rows})


def test_edit_flow_import_actions_export(input_dir, tmp_path):
    from edit_data_manager import EditDataManager, export_excel
    p = _edit_file(input_dir)
    m = EditDataManager(p.name)
    assert (m.schema.column(ColumnRole.STATUS), m.schema.column(ColumnRole.MAPS_URL)) == ("Status", "Map Link")
    recs = m.get_place_records(0)
    assert len(recs) == 50 and all(r.status_value == "REVIEW" for r in recs)
    assert len({r.row_index for r in recs}) == 50                      # tekrar yok
    m.delete_row(recs[0].row_index)
    m.keep_row(recs[1].row_index)
    m.skip_row(recs[2].row_index)
    assert m.undo_action(recs[1].row_index) == "keep"
    m.close()

    m2 = EditDataManager(p.name)                                       # yeniden açılış: içe aktarma yok
    assert m2.get_pending_count() == 49                                # silinen hariç, atlanan geri geldi
    assert m2.get_total_count() == 199
    m2.close()

    out = tmp_path / "out.xlsx"
    assert export_excel(out, p.name) == 199
    ws = openpyxl.load_workbook(out).active
    assert [c.value for c in ws[1]] == ["Status", "Business", "Rating", "Map Link", "Seen"]
    names = [r[1] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert "Biz 0" not in names and len(names) == len(set(names)) == 199
    seen = ws.cell(row=2, column=5).value
    assert hasattr(seen, "year") and seen.year == 2026                 # tarih tipi korundu


def test_edit_flow_requires_status_column(input_dir):
    from data_model import ColumnDetectionError
    from edit_data_manager import EditDataManager
    rows = [("isim", "maps")] + [(f"A{i}", maps_url(i)) for i in range(10)]
    p = write_xlsx(input_dir / "cleanpoi_cevre_nostatus.xlsx", {"S": rows})
    with pytest.raises(ColumnDetectionError):
        EditDataManager(p.name)


def test_legacy_database_without_schema_meta_still_works(input_dir):
    """Eski sürümün oluşturduğu veritabanı (meta'da 'schema' yok) kayıtlı veriden tespit edilir."""
    from edit_data_manager import EditDataManager
    p = _edit_file(input_dir, "cleanpoi_cevre_legacy.xlsx", n=40)
    m = EditDataManager(p.name)
    m.delete_row(0)
    m.close()
    db = config.edit_db_path(p)
    con = sqlite3.connect(db)
    con.execute("DELETE FROM meta WHERE key='schema'")
    con.commit()
    con.close()
    m2 = EditDataManager(p.name)
    assert m2.schema.column(ColumnRole.STATUS) == "Status"
    assert m2.get_pending_count() == 9                                  # 10 REVIEW - 1 silinen
    m2.close()
    con = sqlite3.connect(db)
    assert json.loads(con.execute("SELECT value FROM meta WHERE key='schema'").fetchone()[0])["matches"]
    con.close()
