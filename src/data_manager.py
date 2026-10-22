"""
Proje V1 — Etiketleme Akışı Veri Yönetimi (DataManager)

Excel/CSV dosyasındaki Maps bağlantısı ve isim kolonlarını içerikten tespit eder,
kullanıcının girdiği kategori ve ek alanları dosyaya yazar.

Performans / dayanıklılık:
  • Eski sürüm her etikette tüm .xlsx'i yeniden kaydediyordu (~1 sn, event loop'u
    bloke ediyordu) ve yerinde yazdığı için kesintide dosya bozulabiliyordu.
  • Yeni sürüm: her kayıt önce SQLite günlüğüne (journal) yazılır (ms düzeyi,
    kalıcı), bellekteki sayfaya uygulanır; arka plan thread'i değişiklikleri
    toplayıp dosyayı birkaç saniyede bir ATOMİK olarak (geçici dosya → replace)
    kaydeder. Süreç çökerse açılışta günlükteki yazılmamış kayıtlar yeniden
    uygulanır — veri kaybı olmaz. Dosya Excel'de açık/kilitliyse kayıtlar
    günlükte bekler ve daha sonra tekrar denenir.
"""
from __future__ import annotations

import csv
import json
import logging
import os
import sqlite3
import threading
import time
from pathlib import Path
from typing import Any, Optional

from config import (
    CATEGORY_COLUMN,
    EXTRA_COLUMNS,
    LABEL_JOURNAL_FILE,
    find_label_input_file,
)
from data_model import ColumnDetectionError, ColumnRole, PlaceRecord, SheetSchema, clean_cell
from excel_schema import get_schema, normalize_header

logger = logging.getLogger(__name__)

FLUSH_DEBOUNCE_SEC = 2.0     # son değişiklikten bu kadar sonra kaydet
FLUSH_MAX_DELAY_SEC = 10.0   # sürekli değişiklik olsa bile en geç bu sürede kaydet
FLUSH_RETRY_SEC = 10.0       # dosya kilitliyse tekrar deneme aralığı


class LabelJournal:
    """Excel'e henüz yazılmamış etiketlerin kalıcı günlüğü."""

    def __init__(self, db_file: Path) -> None:
        db_file.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._conn = sqlite3.connect(db_file, check_same_thread=False, isolation_level=None)
        self._conn.execute("PRAGMA journal_mode=WAL")
        self._conn.execute(
            "CREATE TABLE IF NOT EXISTS entries (id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "file TEXT NOT NULL, row_index INTEGER NOT NULL, payload TEXT NOT NULL, ts REAL NOT NULL)")

    def add(self, file: str, row_index: int, values: dict) -> int:
        with self._lock:
            cur = self._conn.execute(
                "INSERT INTO entries (file, row_index, payload, ts) VALUES (?,?,?,?)",
                (file, row_index, json.dumps(values, ensure_ascii=False), time.time()))
            return cur.lastrowid

    def pending(self, file: str) -> list[tuple[int, int, dict]]:
        with self._lock:
            rows = self._conn.execute(
                "SELECT id, row_index, payload FROM entries WHERE file=? ORDER BY id", (file,)).fetchall()
        return [(i, r, json.loads(p)) for i, r, p in rows]

    def max_id(self, file: str) -> int:
        with self._lock:
            return self._conn.execute("SELECT COALESCE(MAX(id),0) FROM entries WHERE file=?", (file,)).fetchone()[0]

    def clear_upto(self, file: str, max_id: int) -> None:
        with self._lock:
            self._conn.execute("DELETE FROM entries WHERE file=? AND id<=?", (file, max_id))

    def close(self) -> None:
        with self._lock:
            self._conn.close()


# ── Tablo arka uçları ────────────────────────────────────────────────────────
class _XlsxTable:
    def __init__(self, path: Path, schema: SheetSchema) -> None:
        import openpyxl
        self.path = path
        self.wb = openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
        self.ws = self.wb[schema.sheet]
        self.header_row = schema.header_row

    def headers(self) -> list[Any]:
        return [c.value for c in self.ws[self.header_row]]

    def add_column(self, name: str) -> int:
        col = self.ws.max_column + 1
        self.ws.cell(row=self.header_row, column=col, value=name)
        return col - 1

    def rows(self):
        return self.ws.iter_rows(min_row=self.header_row + 1, values_only=True)

    def set(self, row_index: int, col_idx: int, value: Any) -> None:
        self.ws.cell(row=self.header_row + 1 + row_index, column=col_idx + 1, value=value)

    def save(self) -> None:
        tmp = self.path.with_name(f"~tmp_{self.path.stem}{self.path.suffix}")
        self.wb.save(tmp)
        try:
            os.replace(tmp, self.path)
        except OSError:
            tmp.unlink(missing_ok=True)
            raise


class _CsvTable:
    def __init__(self, path: Path, schema: SheetSchema) -> None:
        self.path = path
        with open(path, "r", encoding="utf-8-sig", newline="") as fh:
            head = fh.read(8192)
            fh.seek(0)
            try:
                self.dialect = csv.Sniffer().sniff(head, delimiters=",;\t|")
            except csv.Error:
                self.dialect = csv.excel
            all_rows = [list(r) for r in csv.reader(fh, self.dialect)]
        self.pre = all_rows[:max(0, schema.header_row - 1)]
        self.header = all_rows[schema.header_row - 1] if schema.header_row else []
        self.data = all_rows[schema.header_row:]

    def headers(self) -> list[Any]:
        return list(self.header)

    def add_column(self, name: str) -> int:
        self.header.append(name)
        return len(self.header) - 1

    def rows(self):
        for r in self.data:
            yield tuple(v if v != "" else None for v in r)

    def set(self, row_index: int, col_idx: int, value: Any) -> None:
        row = self.data[row_index]
        if len(row) <= col_idx:
            row.extend([""] * (col_idx + 1 - len(row)))
        row[col_idx] = "" if value is None else str(value)

    def save(self) -> None:
        tmp = self.path.with_name(f"~tmp_{self.path.name}")
        with open(tmp, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh, self.dialect)
            w.writerows(self.pre)
            w.writerow(self.header)
            w.writerows(self.data)
        os.replace(tmp, self.path)


class DataManager:
    """
    Excel/CSV veri kaynağını yöneten sınıf.
    - İşlenmemiş (tüm etiket alanları boş) satırları döndürür.
    - Her kayıt anında kalıcıdır (günlük); dosyaya arka planda atomik yazılır.
    """

    def __init__(self, source_file: Optional[str] = None, start_flusher: bool = True,
                 read_only: bool = False) -> None:
        self._file_path: Path = find_label_input_file(source_file)
        self._read_only = read_only   # True: dosyaya hiç yazılmaz (ör. toplu ön-yakalama)
        self._lock = threading.RLock()
        self._journal = LabelJournal(LABEL_JOURNAL_FILE)
        self._journal_key = str(self._file_path.resolve())
        self.schema: SheetSchema = get_schema(self._file_path, ColumnRole.MAPS_URL)
        if self.schema.header_row == 0:
            raise ColumnDetectionError(
                f"'{self._file_path.name}' başlık satırı içermiyor; etiket kolonları eklenemez.")
        self._label_cols: dict[str, int] = {}
        self._records: dict[int, PlaceRecord] = {}
        self._labelled: set[int] = set()
        self._pending_indices: list[int] = []
        self._total = 0

        self._dirty = False
        self._first_dirty = 0.0
        self._last_change = 0.0
        self._next_retry = 0.0
        self._stop = threading.Event()
        self._wake = threading.Event()
        self._flusher: Optional[threading.Thread] = None
        self.flush_stats = {"saves": 0, "save_time": 0.0, "failed_saves": 0}

        t = time.perf_counter()
        self._table = _CsvTable(self._file_path, self.schema) if self._file_path.suffix.lower() == ".csv" \
            else _XlsxTable(self._file_path, self.schema)
        self._load()
        logger.info(f"Dosya yüklendi: {self._file_path.name} — {self._total} satır, "
                    f"{len(self._pending_indices)} bekleyen ({time.perf_counter() - t:.2f} sn)")
        if start_flusher and not read_only:
            self._flusher = threading.Thread(target=self._flush_loop, name="label-flusher", daemon=True)
            self._flusher.start()

    # ── Yükleme ──────────────────────────────────────────────────────────────
    def _load(self) -> None:
        headers = self._table.headers()
        norm_headers = [normalize_header(h) for h in headers]
        for col_name in [CATEGORY_COLUMN, *EXTRA_COLUMNS]:
            idx = next((i for i, h in enumerate(headers) if h == col_name), None)
            if idx is None:
                n = normalize_header(col_name)
                idx = next((i for i, h in enumerate(norm_headers) if h == n), None)
            if idx is None:
                idx = self._table.add_column(col_name)
                self._mark_dirty()
                logger.warning(f"'{col_name}' sütunu bulunamadı, yeni sütun oluşturuldu.")
            self._label_cols[col_name] = idx

        url_i = self.schema.index(ColumnRole.MAPS_URL)
        name_i = self.schema.index(ColumnRole.NAME)
        pid_i = self.schema.index(ColumnRole.PLACE_ID)
        label_idx = list(self._label_cols.values())

        def cell(row: tuple, i: Optional[int]) -> Any:
            return row[i] if i is not None and i < len(row) else None

        for row_index, row in enumerate(self._table.rows()):
            self._total = row_index + 1
            raw_url = cell(row, url_i)
            if not clean_cell(raw_url):
                continue
            self._records[row_index] = PlaceRecord.build(row_index, raw_url, cell(row, name_i), None,
                                                         cell(row, pid_i))
            if any(clean_cell(cell(row, i)) for i in label_idx):
                self._labelled.add(row_index)

        # Çökme kurtarma: dosyaya yazılmamış günlük kayıtlarını uygula
        replay = self._journal.pending(self._journal_key)
        for _, row_index, values in replay:
            self._apply(row_index, values)
        if replay:
            logger.warning(f"{len(replay)} adet dosyaya yazılmamış etiket günlükten geri yüklendi.")
            self._mark_dirty()

        self._pending_indices = sorted(i for i in self._records if i not in self._labelled)

    def _apply(self, row_index: int, values: dict) -> None:
        for col_name, idx in self._label_cols.items():
            self._table.set(row_index, idx, values.get(col_name))
        if any(clean_cell(v) for v in values.values()):
            self._labelled.add(row_index)
        else:
            self._labelled.discard(row_index)

    # ── Kalıcılık ────────────────────────────────────────────────────────────
    def _mark_dirty(self) -> None:
        now = time.monotonic()
        if not self._dirty:
            self._first_dirty = now
        self._dirty = True
        self._last_change = now
        self._wake.set()

    def _flush_loop(self) -> None:
        while not self._stop.is_set():
            self._wake.wait(timeout=1.0)
            self._wake.clear()
            if not self._dirty:
                continue
            now = time.monotonic()
            due = (now - self._last_change >= FLUSH_DEBOUNCE_SEC) or (now - self._first_dirty >= FLUSH_MAX_DELAY_SEC)
            if due and now >= self._next_retry:
                self.flush()

    def flush(self) -> bool:
        """Bekleyen değişiklikleri dosyaya atomik olarak yazar. Başarılıysa True."""
        with self._lock:
            if not self._dirty or self._read_only:
                return True
            upto = self._journal.max_id(self._journal_key)
            t = time.perf_counter()
            try:
                self._table.save()
            except OSError as e:
                self.flush_stats["failed_saves"] += 1
                self._next_retry = time.monotonic() + FLUSH_RETRY_SEC
                logger.warning(f"Dosya kaydedilemedi (Excel'de açık olabilir), kayıtlar günlükte güvende; "
                               f"{FLUSH_RETRY_SEC:.0f} sn sonra tekrar denenecek: {e}")
                return False
            self._journal.clear_upto(self._journal_key, upto)
            self._dirty = False
            self.flush_stats["saves"] += 1
            self.flush_stats["save_time"] += time.perf_counter() - t
            return True

    def close(self) -> None:
        self._stop.set()
        self._wake.set()
        if self._flusher:
            self._flusher.join(timeout=5)
        self.flush()
        self._journal.close()

    # ── Public API ───────────────────────────────────────────────────────────
    def get_pending_count(self) -> int:
        """İşlenmemiş satır sayısını döndürür."""
        with self._lock:
            return sum(1 for i in self._records if i not in self._labelled)

    def get_total_count(self) -> int:
        """Toplam veri satırı sayısı."""
        return self._total

    def get_record(self, position: int) -> Optional[dict]:
        """Belirli bir konumdaki (0-indexed) işlenmemiş kaydı döndürür."""
        if position >= len(self._pending_indices):
            return None
        rec = self._records[self._pending_indices[position]]
        d = rec.to_dict()
        d["position"], d["total"] = position, len(self._pending_indices)
        return d

    def get_urls_from(self, start_position: int, count: int) -> list[dict]:
        """start_position'dan itibaren count adet kayıt döndürür (buffer için)."""
        end = min(start_position + count, len(self._pending_indices))
        return [r for r in (self.get_record(p) for p in range(start_position, end)) if r]

    def get_place_records(self, start_index: int = 0) -> list[PlaceRecord]:
        """Veri indeksinden itibaren işlenmemiş kayıtları, oturum sırasıyla döndürür."""
        if start_index < 0:
            raise ValueError("Başlangıç indeksi 0 veya daha büyük olmalıdır.")
        with self._lock:
            idx = [i for i in self._pending_indices if i >= start_index and i not in self._labelled]
        out = []
        for position, i in enumerate(idx):
            src = self._records[i]
            out.append(PlaceRecord(**{**src.__dict__, "position": position, "total": len(idx)}))
        return out

    def get_pending_records_from_index(self, start_index: int) -> list[dict]:
        return [r.to_dict() for r in self.get_place_records(start_index)]

    def save_category(self, row_index: int, category: str, extra_fields: Optional[dict] = None) -> None:
        """
        Satırın kategori ve ek alanlarını yazar (boş bırakılan eski değerler temizlenir).
        Günlüğe yazıldığı anda kalıcıdır; dosyaya arka planda kaydedilir.
        """
        if self._read_only:
            raise PermissionError("DataManager salt-okunur modda açıldı.")
        if row_index < 0 or row_index >= self._total:
            raise IndexError(f"Geçersiz satır: {row_index}")
        values = {CATEGORY_COLUMN: category or None}
        for col_name in EXTRA_COLUMNS:
            values[col_name] = (extra_fields or {}).get(col_name) or None
        with self._lock:
            self._journal.add(self._journal_key, row_index, values)
            self._apply(row_index, values)
            self._mark_dirty()
        logger.info(f"Kaydedildi — satır {row_index}: kategori='{category}', ek alanlar={extra_fields}")
